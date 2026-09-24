#!/usr/bin/env python3
"""Export all completed simulator trades to Excel and recompute C4-C7 at SellTime.

Run from the project container, e.g. telegram-notifier, where the database volume
and current telegram_notifications.yaml are mounted.
"""
from __future__ import annotations

import json
import math
import os
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from server.database import Instrument, Measurement, SessionLocal, SimulationTrade
from shared.trading_decisions import evaluate_sell_history, trading_window_info

CONFIG_PATH = Path(os.getenv("TELEGRAM_NOTIFICATIONS_CONFIG", "/app/server/telegram_notifications.yaml"))
OUTPUT_PATH = Path(os.getenv("SIM_TRADES_XLSX", "/tmp/simulation_trades.xlsx"))
LOCAL_TZ = ZoneInfo(os.getenv("EXPORT_TIMEZONE", "Europe/Berlin"))


def market_region(ticker: str, asset_type: str | None, isin: str | None) -> str:
    if str(asset_type or "").lower() == "crypto":
        return "CRYPTO"
    isin = str(isin or "").upper()
    if isin.startswith("DE"):
        return "DE"
    # The stock collector is Massive/Polygon US unless explicitly identifiable as DE.
    return "US"


def as_utc(value):
    ts = pd.to_datetime(value, utc=True, errors="coerce")
    return ts


def local_naive(value):
    ts = as_utc(value)
    if pd.isna(ts):
        return None
    # Excel does not support timezone-aware datetimes.
    return ts.tz_convert(LOCAL_TZ).to_pydatetime().replace(tzinfo=None)


def parse_measurement(row: Measurement):
    try:
        values = json.loads(row.measurements_json or "{}")
    except Exception:
        values = {}
    try:
        metadata = json.loads(row.metadata_json or "{}")
    except Exception:
        metadata = {}
    ticker = str(metadata.get("ticker") or row.system or "").strip().upper()
    asset_type = str(metadata.get("asset_type") or ("crypto" if row.system == "crypto" else "stock"))
    close = pd.to_numeric(values.get("close"), errors="coerce")
    return ticker, asset_type, close


def regular_close_remaining_minutes(action_time, market_config):
    window = trading_window_info(action_time, market_config, "sell")
    if not window.is_open:
        return window, None
    regular_close = market_config.get("regular_close")
    if not regular_close:
        return window, None
    try:
        tz_name = str(market_config.get("timezone") or "UTC")
        local = as_utc(action_time).tz_convert(tz_name)
        hh, mm = [int(x) for x in str(regular_close).split(":", 1)]
        close_local = pd.Timestamp(
            datetime(local.year, local.month, local.day, hh, mm), tz=tz_name
        )
        remaining = (close_local - local).total_seconds() / 60.0
        return window, remaining
    except Exception:
        return window, None


def main():
    if not CONFIG_PATH.exists():
        raise SystemExit(f"Config not found: {CONFIG_PATH}")
    config = yaml.safe_load(CONFIG_PATH.read_text(encoding="utf-8")) or {}
    sell_cfg = config.get("sell") or {}
    windows = config.get("trading_windows") or {}
    phases = config.get("trading_phases") or {}

    movement = float(sell_cfg.get("movement_percent", 1.1))
    c5_hours = float(sell_cfg.get("c5_hours", 24.0))
    c6_enabled = bool(sell_cfg.get("c6_enabled", True))
    c6_close_minutes = float(sell_cfg.get("c6_close_minutes", 30.0))
    c6_min_gain = float(sell_cfg.get("c6_min_gain_percent", 2.0))
    c7_max_gain = float(sell_cfg.get("c7_max_gain_percent", 5.0))

    with SessionLocal() as db:
        trades = db.scalars(
            select(SimulationTrade)
            .where(
                SimulationTrade.buy_telegram_sent.is_(True),
                SimulationTrade.sell_time.is_not(None),
            )
            .order_by(SimulationTrade.buy_time.asc(), SimulationTrade.id.asc())
        ).all()

        instruments = {r.ticker.upper(): r for r in db.scalars(select(Instrument)).all()}

        # Read market history once. C4/C5 need the full retained history between BuyTime and SellTime.
        mrows = db.scalars(select(Measurement).order_by(Measurement.timestamp.asc(), Measurement.id.asc())).all()

    history = {}
    asset_types = {}
    for m in mrows:
        ticker, asset_type, close = parse_measurement(m)
        if not ticker or pd.isna(close):
            continue
        history.setdefault(ticker, []).append({
            "timestamp": as_utc(m.timestamp),
            "close": float(close),
            "id": m.id,
        })
        asset_types[ticker] = asset_type

    history_df = {
        ticker: pd.DataFrame(rows).sort_values(["timestamp", "id"])
        for ticker, rows in history.items()
    }

    export_rows = []
    for trade in trades:
        ticker = trade.ticker.upper()
        buy_time = as_utc(trade.buy_time)
        sell_time = as_utc(trade.sell_time)
        inst = instruments.get(ticker)
        asset_type = asset_types.get(ticker) or (inst.asset_type if inst else "stock")
        region = market_region(ticker, asset_type, inst.isin if inst else None)
        market_cfg = windows.get(region) or {}

        ticker_history = history_df.get(ticker, pd.DataFrame(columns=["timestamp", "close", "id"]))
        ticker_history = ticker_history[ticker_history["timestamp"] <= sell_time].copy()

        sell_price = pd.to_numeric(trade.sell_price, errors="coerce")
        buy_price = pd.to_numeric(trade.buy_price, errors="coerce")

        c4 = c5 = c6 = c7 = False
        if not ticker_history.empty and pd.notna(sell_price) and float(sell_price) > 0:
            decision = evaluate_sell_history(
                ticker_df=ticker_history,
                latest_time=sell_time,
                current_price=float(sell_price),
                movement_percent=movement,
                c5_hours=c5_hours,
                init_time=buy_time,
                market_region=region,
                market_config=market_cfg or None,
                phase_config=phases.get(region) or {},
            )
            c4 = bool(decision.c4_satisfied)
            c5 = bool(decision.c5_satisfied)

            if market_cfg:
                window, remaining = regular_close_remaining_minutes(sell_time, market_cfg)
                gain = None
                if pd.notna(buy_price) and float(buy_price) > 0:
                    gain = (float(sell_price) / float(buy_price) - 1.0) * 100.0
                c6 = bool(
                    c6_enabled and not c4 and not c5 and window.is_open
                    and remaining is not None and 0.0 <= remaining <= c6_close_minutes
                    and gain is not None and gain < c6_min_gain
                )
                c7 = bool(
                    not c4 and not c5 and window.is_open
                    and remaining is not None and 0.0 <= remaining <= c6_close_minutes
                    and gain is not None and gain > c7_max_gain
                )

        diff_pct = pd.to_numeric(trade.relative_difference, errors="coerce")
        if pd.isna(diff_pct) and pd.notna(buy_price) and pd.notna(sell_price) and float(buy_price) != 0:
            diff_pct = (float(sell_price) / float(buy_price) - 1.0) * 100.0

        # Match simulator/Logs sizing: smallest whole-share quantity covering EUR 10,000.
        buy_eur = pd.to_numeric(trade.buy_price_eur, errors="coerce")
        sell_eur = pd.to_numeric(trade.sell_price_eur, errors="coerce")
        diff_eur = None
        if pd.notna(buy_eur) and float(buy_eur) > 0:
            qty = math.ceil(10000.0 / float(buy_eur))
            if pd.notna(sell_eur):
                diff_eur = (float(sell_eur) - float(buy_eur)) * qty
            elif pd.notna(diff_pct):
                diff_eur = float(buy_eur) * qty * float(diff_pct) / 100.0

        export_rows.append({
            "Ticker": ticker,
            "InitTime": local_naive(buy_time),
            "EndTime": local_naive(sell_time),
            "DiffSellPrice%": None if pd.isna(diff_pct) else float(diff_pct) / 100.0,
            "DiffSellPrice": diff_eur,
            "SimSellC4": c4,
            "SimSellC5": c5,
            "SimSellC6": c6,
            "SimSellC7": c7,
        })

    wb = Workbook()
    ws = wb.active
    ws.title = "Sim Trades"
    headers = ["Ticker", "InitTime", "EndTime", "DiffSellPrice%", "DiffSellPrice", "SimSellC4", "SimSellC5", "SimSellC6", "SimSellC7"]
    ws.append(headers)
    for row in export_rows:
        ws.append([row[h] for h in headers])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = [12, 20, 20, 17, 17, 13, 13, 13, 13]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    for row in range(2, ws.max_row + 1):
        ws.cell(row, 2).number_format = "yyyy-mm-dd hh:mm"
        ws.cell(row, 3).number_format = "yyyy-mm-dd hh:mm"
        ws.cell(row, 4).number_format = '+0.00%;[Red]-0.00%;-'
        ws.cell(row, 5).number_format = '€#,##0.00;[Red](€#,##0.00);-'

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)
    print(f"Exported {len(export_rows)} completed simulator trades to {OUTPUT_PATH}")
    print(f"C4/C5 settings: movement={movement}% c5_hours={c5_hours}")
    print(f"C6/C7 settings: close_minutes={c6_close_minutes}, C6<{c6_min_gain}%, C7>{c7_max_gain}%")


if __name__ == "__main__":
    main()
